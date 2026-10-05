"""Build an evidence-labelled, non-overwriting Word update from the v2 report."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
RUNS = ROOT / "dataset_crack" / "runs" / "pipe_proto"
SOURCE = ROOT / "BFY_ModelComp_report_v2.docx"
OUTPUT = ROOT / "BFY_ModelComp2_report_updated.docx"


def best_yolo(run: str) -> dict[str, float | int]:
    with (RUNS / run / "results.csv").open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    row = max(rows, key=lambda item: float(item["metrics/mAP50-95(B)"]))
    return {
        "epoch": int(row["epoch"]),
        "logged_epochs": len(rows),
        "precision": float(row["metrics/precision(B)"]),
        "recall": float(row["metrics/recall(B)"]),
        "map50": float(row["metrics/mAP50(B)"]),
        "map5095": float(row["metrics/mAP50-95(B)"]),
    }


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def pct(value: float) -> str:
    return f"{100 * value:.2f}%"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc, headers: list[str], rows: list[list[str]]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = True
    for idx, title in enumerate(headers):
        cell = result.rows[0].cells[idx]
        cell.text = title
        set_cell_shading(cell, "17365D")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, values in enumerate(rows):
        cells = result.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = str(value)
            if index % 2:
                set_cell_shading(cells[col], "EAF0F7")
    for row in result.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()


def add_text(doc, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)


def append_legacy_record(doc, source: Document) -> None:
    """Rebuild text/tables cleanly; the old DOCX lacks table grids."""
    doc.add_heading("Historical research record (29 September 2026 snapshot)", level=1)
    add_text(
        doc,
        "The following archived sections preserve the previous dataset, architecture, "
        "literature and experiment notes. Earlier conclusions were written before the "
        "RF-DETR run above and are superseded wherever they conflict with the October update. "
        "The YOLO11n rows are illustrative placeholders, not actual runs.",
    )
    for child in source._body._body.iterchildren():
        if child.tag == qn("w:p"):
            words = "".join(node.text or "" for node in child.iter(qn("w:t")))
            if not words.strip():
                continue
            if words.startswith("RoadEye Sewer Defect Detection") or words in {
                "Consolidated research record",
                "This report consolidates the previous comparison sheet and the v2-to-v5 dataset documentation. Measured values are listed separately from illustrative placeholders.",
            }:
                continue
            # Keep the original section hierarchy without carrying over stale formatting.
            if words[:2].rstrip(".").isdigit() and ". " in words[:5]:
                doc.add_heading(words, level=2)
            elif words in {"Unseen test split comparison", "Protocol and reproducibility", "Evidence files"}:
                doc.add_heading(words, level=3)
            else:
                add_text(doc, words)
        elif child.tag == qn("w:tbl"):
            rows = []
            for tr in child.findall(qn("w:tr")):
                values = []
                for tc in tr.findall(qn("w:tc")):
                    values.append(" ".join(node.text or "" for node in tc.iter(qn("w:t"))))
                if values:
                    rows.append(values)
            if rows:
                width = len(rows[0])
                table(doc, rows[0], [row[:width] + [""] * max(0, width - len(row)) for row in rows[1:]])


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Refusing to overwrite existing report: {OUTPUT}")
    if not SOURCE.exists():
        raise SystemExit(f"Missing previous report: {SOURCE}")

    baseline = best_yolo("yolo26s_v5_cctv")
    simam = best_yolo("yolo26s_v5_clahe_simam")
    cbam = best_yolo("yolo26s_v5_clahe_cbam")
    se = best_yolo("yolo26s_v5_clahe_se")
    parallel = best_yolo("yolo26s_v5_parallel_simam_cbam_se")
    v4 = best_yolo("yolo26n_v4_pipe_cctv_clahe-3")
    v6 = best_yolo("yolo26s_v6_normal_error_filtered_diagnostic_seed0_20260930_190356_a5b23911")
    rf40_dir = RUNS / "rfdetr_nano_v5_normal_512_seed0_20261002_193202_f6bd13"
    rf50_dir = RUNS / "rfdetr_nano_v5_normal_512_seed0_20261003_193042_66b2b7"
    rf40 = read_json(rf40_dir / "best_validation_metrics.json")
    rf50 = read_json(rf50_dir / "best_validation_metrics.json")
    rf_config = read_json(rf50_dir / "training_config.json")
    provenance = read_json(rf50_dir / "dataset_provenance.json")
    biased = read_json(
        RUNS
        / "yolo26s_v6_normal_error_filtered_diagnostic_seed0_20260929_184244_3e235d15_filtered_test"
        / "diagnostic_summary.json"
    )
    full_test = read_json(
        RUNS
        / "yolo26s_v6_normal_error_filtered_diagnostic_seed0_20260929_184244_3e235d15_full_v5_test"
        / "evaluation_summary.json"
    )

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin = section.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    doc.styles["Title"].font.color.rgb = RGBColor(23, 54, 93)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("RoadEye | Model Comparison and Research Record")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Updated 4 October 2026 · Measured results, incomplete runs and limitations")

    doc.add_heading("Executive finding", level=1)
    add_text(
        doc,
        f"Yes: RF-DETR Nano after 50 total training epochs reached {pct(rf50['val/mAP_50'])} "
        f"validation mAP@0.50 using its saved best checkpoint. Its validation mAP@0.50:0.95 "
        f"was {pct(rf50['val/mAP_50_95'])}; precision {pct(rf50['val/precision'])}, "
        f"recall {pct(rf50['val/recall'])}, and F1 {pct(rf50['val/F1'])}. "
        "This is a validation result, not 90% accuracy on the untouched test split. "
        "No RF-DETR test-set result was found in this run (run_test=false).",
    )
    add_text(
        doc,
        f"The previous 40-epoch checkpoint scored {pct(rf40['val/mAP_50'])} mAP50 and "
        f"{pct(rf40['val/mAP_50_95'])} mAP50-95. The additional 10 epochs raised those "
        f"saved-checkpoint validation scores by {(rf50['val/mAP_50']-rf40['val/mAP_50'])*100:.2f} "
        f"and {(rf50['val/mAP_50_95']-rf40['val/mAP_50_95'])*100:.2f} percentage points, respectively.",
    )

    doc.add_heading("October 2026 experiment ledger", level=1)
    add_text(
        doc,
        "Unless labelled otherwise, YOLO figures below are the single validation CSV row with "
        "the highest mAP50-95 in that run. E3/E5 values are from the prior report's explicit "
        "post-training validation summaries. Architecture failures mean no measured gain over "
        "the matched v5 baseline, not a failed training process.",
    )
    ledger = [
        ["v4 YOLO26n + CLAHE", "v4 CCTV", str(v4['logged_epochs']), pct(v4['map50']), pct(v4['map5095']), "Historical; different dataset"],
        ["YOLO26s baseline", "v5 normal", "80", pct(baseline['map50']), pct(baseline['map5095']), "Best YOLO validation comparator"],
        ["YOLO26s + SimAM", "v5 CLAHE", "80", pct(simam['map50']), pct(simam['map5095']), "No clear gain; CLAHE also differs"],
        ["YOLO26s + CBAM", "v5 CLAHE", "80", pct(cbam['map50']), pct(cbam['map5095']), "mAP50 near tie; lower mAP50-95"],
        ["YOLO26s + SE", "v5 CLAHE", "80", pct(se['map50']), pct(se['map5095']), "No gain"],
        ["Parallel SimAM+CBAM+SE", "v5 normal", "80", pct(parallel['map50']), pct(parallel['map5095']), "No gain"],
        ["E1 P2 head", "v5 normal", "—", "—", "—", "Pending; no measured result found"],
        ["E2 weighted PAN", "v5 normal", "—", "—", "—", "Pending; no measured result found"],
        ["E3 P2 + weighted PAN", "v5 normal", "80", "82.19%", "65.01%", "Completed; near baseline"],
        ["E4 YOLO26m", "v5 normal", "—", "—", "—", "Pending; no measured result found"],
        ["E5 directional P3 block", "v5 normal", "80", "82.01%", "63.82%", "Completed; no gain"],
        ["E6 E3 + SimAM", "v5 normal", "4/80", "—", "—", "OOM in epoch 5; inconclusive"],
        ["v6 error-filtered YOLO26s", "v6 filtered", "80", pct(v6['map50']), pct(v6['map5095']), "Degraded; diagnostic subset"],
        ["RF-DETR Nano", "v5 normal", "40", pct(rf40['val/mAP_50']), pct(rf40['val/mAP_50_95']), "Completed; below 90% mAP50"],
        ["RF-DETR Nano resumed", "v5 normal", "50 total", pct(rf50['val/mAP_50']), pct(rf50['val/mAP_50_95']), "Crossed 90% validation mAP50"],
    ]
    table(doc, ["Experiment", "Data", "Epochs", "Val mAP50", "Val mAP50-95", "Outcome"], ledger)

    doc.add_heading("Why the v6 high number is not the passing result", level=1)
    add_text(
        doc,
        f"The model-error-selected filtered test reported {pct(biased['mAP50'])} mAP50 "
        f"and {pct(biased['mAP50_95'])} mAP50-95, but excluded hard test images based on "
        "prior predictions. That is selection-biased and cannot be reported as a held-out "
        f"benchmark. The same checkpoint scored {pct(full_test['mAP50'])} mAP50 and "
        f"{pct(full_test['mAP50_95'])} mAP50-95 on the unchanged full v5 test split. "
        f"After removing model-selected FP images and 80% of FN images from v6 training "
        f"({v6['logged_epochs']} logged epochs), validation fell to {pct(v6['map50'])} mAP50 "
        f"and {pct(v6['map5095'])} mAP50-95. This supports retaining hard examples and "
        "auditing annotations rather than deleting images solely because a model misses them.",
    )

    doc.add_heading("RF-DETR protocol and fair comparison", level=1)
    config = rf_config["train_config"]
    model = rf_config["model_config"]
    add_text(
        doc,
        f"Architecture: {model['model_name']} (pretrained), {model['resolution']}-pixel input; "
        f"{config['batch_size']} image per GPU step with gradient accumulation "
        f"{config['grad_accum_steps']} (effective batch {config['batch_size']*config['grad_accum_steps']}), "
        f"EMA enabled, seed {config['seed']}, local RTX 4070 Laptop GPU. "
        f"Normal v5 has {provenance['counts']['train']:,} train, "
        f"{provenance['counts']['valid']:,} validation and {provenance['counts']['test']:,} "
        "test images. The RF-DETR overlay kept all images and split assignments, omitted "
        f"{provenance['removed_label_row_count']} invalid zero-area box rows "
        "(three train, one validation), and did not use offline CLAHE.",
    )
    add_text(
        doc,
        f"The best YOLO26s v5 normal validation row was {pct(baseline['map50'])} mAP50 "
        f"and {pct(baseline['map5095'])} mAP50-95 (epoch {baseline['epoch']}/80). "
        f"RF-DETR's recorded validation advantage is "
        f"{(rf50['val/mAP_50']-baseline['map50'])*100:.2f} points mAP50 and "
        f"{(rf50['val/mAP_50_95']-baseline['map5095'])*100:.2f} points mAP50-95. "
        "Treat this as indicative, not a strict architecture-only benchmark: input size "
        "(512 vs 640), evaluation backends/maximum detections and one invalid validation "
        "box differ. Precision/recall are also threshold-dependent. Re-evaluate both "
        "checkpoints on the identical untouched test protocol before making a paper claim.",
    )
    add_text(
        doc,
        "The log's 'Val (Epoch 1/100)' after training is a standalone best-checkpoint "
        "validation pass. It does not indicate another 100 training epochs. The 50th "
        "epoch EMA table (90.18% mAP50, 76.42% mAP50-95) is distinct from the saved "
        "best-checkpoint validation result (90.65%, 76.51%) reported above.",
    )

    doc.add_heading("Class-wise limitation", level=1)
    add_text(doc, "RF-DETR epoch-50 EMA validation AP50-95 (not a test result):")
    table(doc, ["Class", "AP50-95", "Interpretation"], [
        ["crack", "53.07%", "Weakest; missed/thin cracks and box localization remain priorities"],
        ["corrosion_rust", "88.91%", "Strong"],
        ["sediment_deposit", "70.79%", "Moderate"],
        ["root_intrusion", "91.83%", "Strong"],
        ["joint_defect", "77.52%", "Moderate"],
    ])
    add_text(
        doc,
        "The overall 90.65% mAP50 should not be paraphrased as every class exceeding "
        "90%, nor as 90% exact-box accuracy. Crack localization remains the main "
        "per-class limitation in the reported epoch-50 table.",
    )

    doc.add_heading("Evidence and remaining work", level=1)
    for item in [
        "RF-DETR 50 epochs: dataset_crack/runs/pipe_proto/rfdetr_nano_v5_normal_512_seed0_20261003_193042_66b2b7/best_validation_metrics.json; training_config.json; dataset_provenance.json",
        "RF-DETR 40 epochs: dataset_crack/runs/pipe_proto/rfdetr_nano_v5_normal_512_seed0_20261002_193202_f6bd13/best_validation_metrics.json",
        "YOLO validation: dataset_crack/runs/pipe_proto/<run_name>/results.csv, with run names shown in the ledger",
        "v6 bias check: dataset_crack/runs/pipe_proto/yolo26s_v6_normal_error_filtered_diagnostic_seed0_20260929_184244_3e235d15_{filtered_test,full_v5_test}/diagnostic_summary.json or evaluation_summary.json",
        "E3/E5 and E6 status: archived September report section 8 and its evidence-file list below",
        "Next: freeze model choice, run RF-DETR once on the untouched v5 test split, log per-class AP and the exact evaluation settings; do not tune on test outcomes.",
    ]:
        add_text(doc, "• " + item)

    append_legacy_record(doc, Document(SOURCE))
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
