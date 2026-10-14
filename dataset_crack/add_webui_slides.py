"""Add Web UI demonstration slides without replacing edits in the existing deck.

Run capture_webui_for_presentation.py first while the server is on port 8001.
This script reads RoadEye_Crack_Detection_Presentation.pptx and writes a new file.
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "RoadEye_Crack_Detection_Presentation.pptx"
OUTPUT = ROOT / "RoadEye_Crack_Detection_Presentation_with_WebUI.pptx"
IMAGES = ROOT / "dataset_crack" / "presentation_assets"
INK = RGBColor(38, 42, 48)
MUTED = RGBColor(91, 98, 106)


def text(slide, value, x, y, w, h, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def title(slide, label, number):
    text(slide, label, .72, .44, 11.8, .62, size=29, bold=True)
    text(slide, str(number), 12.25, 7.07, .4, .25, size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def add_image(slide, path, x, y, w, h):
    with Image.open(path) as image:
        ratio = image.width / image.height
    target = w / h
    if ratio > target:
        width, height = w, w / ratio
    else:
        height, width = h, h * ratio
    slide.shapes.add_picture(str(path), Inches(x + (w-width)/2), Inches(y + (h-height)/2), width=Inches(width), height=Inches(height))


def bullet(slide, lines, x, y, w, h, size=17):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = "•  " + line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = INK
        paragraph.space_after = Pt(15)


def move_before_last(prs, slides):
    """Move newly appended slides immediately before the existing references."""
    ids = prs.slides._sldIdLst
    existing_reference = ids[-len(slides)-1]
    reference_index = list(ids).index(existing_reference)
    for slide in slides:
        element = next(item for item in ids if item.id == slide.slide_id)
        ids.remove(element)
        ids.insert(reference_index, element)
        reference_index += 1


def main():
    for name in ("webui_image_inference.png", "webui_video_upload.png", "webui_live_camera.png", "webui_model_register.png"):
        if not (IMAGES / name).exists():
            raise FileNotFoundError(f"Capture the local UI first: {IMAGES / name}")
    prs = Presentation(SOURCE)
    original_count = len(prs.slides)
    if original_count < 2:
        raise ValueError("The edited presentation does not have a references slide.")
    blank = prs.slide_layouts[6]
    added = []

    slide = prs.slides.add_slide(blank)
    added.append(slide)
    title(slide, "Web UI · image inspection", original_count)
    add_image(slide, IMAGES / "webui_image_inference.png", .62, 1.33, 9.05, 4.75)
    bullet(slide, ["Choose a local model and confidence threshold.", "Drop an image or select one of five test examples.", "Review annotated boxes, counts and model latency; download the result."], 9.82, 1.55, 2.85, 4.4, size=16)
    text(slide, "Actual local example inference shown. These boxes are predictions, not ground-truth labels.", .73, 6.31, 11.9, .39, size=14, color=MUTED)

    slide = prs.slides.add_slide(blank)
    added.append(slide)
    title(slide, "Web UI · video and live camera inputs", original_count+1)
    text(slide, "Uploaded video", .72, 1.24, 5.8, .35, size=18, bold=True)
    text(slide, "Live feed", 6.78, 1.24, 5.8, .35, size=18, bold=True)
    add_image(slide, IMAGES / "webui_video_upload.png", .69, 1.66, 5.92, 3.08)
    add_image(slide, IMAGES / "webui_live_camera.png", 6.72, 1.66, 5.92, 3.08)
    bullet(slide, ["Video file upload, frame stride, progress and annotated MP4 download.", "Browser webcam, USB camera and opt-in RTSP/HTTP stream sources.", "Live view reports detections, latency and delivered-frame FPS."], .77, 5.0, 11.65, 1.46, size=17)
    text(slide, "Input controls shown; the viewport retains the prior still-image result. USB/IP capture requires a field test.", .76, 6.55, 11.6, .34, size=13, color=MUTED)

    slide = prs.slides.add_slide(blank)
    added.append(slide)
    title(slide, "Web UI · checkpoint register", original_count+2)
    add_image(slide, IMAGES / "webui_model_register.png", .69, 1.28, 11.93, 5.19)
    text(slide, "Local checkpoints are listed with availability, parameter counts and saved validation/test metrics. Different datasets and evaluators are not a head-to-head ranking.", .74, 6.52, 11.73, .40, size=14, color=MUTED)

    move_before_last(prs, added)
    references = prs.slides[-1]
    for shape in references.shapes:
        if shape.has_text_frame and shape.text.strip().isdigit() and len(shape.text.strip()) <= 2:
            shape.text = str(len(prs.slides))
    prs.save(OUTPUT)
    print(f"Preserved {original_count} edited slides; added {len(added)} Web UI slides. Saved {len(prs.slides)} slides to {OUTPUT}")


if __name__ == "__main__":
    main()
